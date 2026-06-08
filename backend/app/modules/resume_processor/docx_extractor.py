import docx
import io

def extract_text_from_docx(file_data: bytes) -> dict:
    # Convert bytes to file object
    file_obj = io.BytesIO(file_data)
    
    try:
        # Open the word document
        doc = docx.Document(file_obj)
        
        text = ""
        
        # Extract text from all paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        
        # Extract text from tables too
        # Some resumes use tables for layout
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
        
        return {
            "text": text,
            "success": True,
            "method": "python-docx"
        }
    
    except Exception as e:
        return {
            "text": "",
            "success": False,
            "method": "python-docx",
            "error": str(e)
        }